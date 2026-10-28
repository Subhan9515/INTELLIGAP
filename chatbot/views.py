import base64
import json
import os
import tempfile
import mimetypes
from pathlib import Path

from django.http import JsonResponse
from django.shortcuts import render
from dotenv import load_dotenv
from google import genai

from .models import ChatMessage


# ============================================================
# LOAD API KEY
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

load_dotenv(BASE_DIR / ".env")


# ============================================================
# GEMINI MODEL
# ============================================================

GEMINI_MODEL = "gemini-3.5-flash-lite"


# ============================================================
# CHAT OWNER
# ============================================================

def get_chat_owner(request):

    if request.user.is_authenticated:
        return f"user_{request.user.pk}"

    if not request.session.session_key:
        request.session.create()

    return f"session_{request.session.session_key}"


# ============================================================
# FILE TYPES
# ============================================================

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
    ".heic",
    ".heif",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".csv",
}

PDF_EXTENSIONS = {
    ".pdf",
}

DOCX_EXTENSIONS = {
    ".docx",
}

PPTX_EXTENSIONS = {
    ".pptx",
}

XLSX_EXTENSIONS = {
    ".xlsx",
}


# ============================================================
# MIME TYPE
# ============================================================

def get_mime_type(uploaded_file):

    mime_type = uploaded_file.content_type

    if mime_type:
        return mime_type

    mime_type = mimetypes.guess_type(
        uploaded_file.name
    )[0]

    return mime_type or "application/octet-stream"


# ============================================================
# FILE CATEGORY
# ============================================================

def get_file_category(uploaded_file):

    extension = Path(
        uploaded_file.name
    ).suffix.lower()

    mime_type = get_mime_type(
        uploaded_file
    )

    # IMAGE
    if extension in IMAGE_EXTENSIONS:
        return "image", mime_type

    # PDF
    if extension in PDF_EXTENSIONS:
        return "pdf", "application/pdf"

    # TEXT
    if extension in TEXT_EXTENSIONS:
        return "text", mime_type

    # WORD
    if extension in DOCX_EXTENSIONS:
        return (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    # POWERPOINT
    if extension in PPTX_EXTENSIONS:
        return (
            "pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

    # EXCEL
    if extension in XLSX_EXTENSIONS:
        return (
            "xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    return None, mime_type


# ============================================================
# EXTRACT TEXT FILE
# ============================================================

def extract_text_file(uploaded_file):

    data = uploaded_file.read()

    try:

        return data.decode(
            "utf-8"
        )

    except UnicodeDecodeError:

        return data.decode(
            "latin-1",
            errors="ignore"
        )


# ============================================================
# EXTRACT DOCX
# ============================================================

def extract_docx(uploaded_file):

    from docx import Document

    document = Document(
        uploaded_file
    )

    parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                row_text.append(
                    cell_text
                )

            parts.append(
                " | ".join(row_text)
            )

    return "\n".join(parts)


# ============================================================
# EXTRACT PPTX
# ============================================================

def extract_pptx(uploaded_file):

    from pptx import Presentation

    presentation = Presentation(
        uploaded_file
    )

    slides_text = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_parts = []

        for shape in slide.shapes:

            if hasattr(
                shape,
                "text"
            ):

                text = shape.text.strip()

                if text:
                    slide_parts.append(
                        text
                    )

        if slide_parts:

            slides_text.append(
                f"Slide {slide_number}\n"
                + "\n".join(slide_parts)
            )

    return "\n\n".join(
        slides_text
    )


# ============================================================
# EXTRACT XLSX
# ============================================================

def extract_xlsx(uploaded_file):

    from openpyxl import load_workbook

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True
    )

    output = []

    for worksheet in workbook.worksheets:

        output.append(
            f"Sheet: {worksheet.title}"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is not None:

                    values.append(
                        str(value)
                    )

                else:

                    values.append("")
            
            if any(values):

                output.append(
                    " | ".join(values)
                )

    workbook.close()

    return "\n".join(output)


# ============================================================
# SAVE TEMPORARY FILE
# ============================================================

def save_uploaded_file(uploaded_file):

    suffix = Path(
        uploaded_file.name
    ).suffix

    temporary_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=suffix
    )

    try:

        for chunk in uploaded_file.chunks():

            temporary_file.write(
                chunk
            )

    finally:

        temporary_file.close()

    return temporary_file.name


# ============================================================
# UPLOAD PDF TO GEMINI
# ============================================================

def upload_pdf_to_gemini(
    client,
    uploaded_file
):

    temporary_path = None

    try:

        temporary_path = save_uploaded_file(
            uploaded_file
        )

        gemini_file = client.files.upload(
            file=temporary_path,
            config={
                "mime_type": "application/pdf"
            }
        )

        return gemini_file

    finally:

        if temporary_path:

            try:

                os.remove(
                    temporary_path
                )

            except OSError:

                pass


# ============================================================
# CHAT
# ============================================================

def chat(request):

    owner = get_chat_owner(
        request
    )

    # ========================================================
    # GET
    # ========================================================

    if request.method == "GET":

        messages = ChatMessage.objects.filter(
            session_key=owner
        ).order_by(
            "created_at",
            "id"
        )

        history = [
            {
                "role": item.role,
                "message": item.message,
            }
            for item in messages
        ]

        return render(
            request,
            "chatbot/chat.html",
            {
                "chat_history": history
            }
        )

    # ========================================================
    # POST
    # ========================================================

    if request.method == "POST":

        try:

            # =================================================
            # GET FILE
            # =================================================

            uploaded_file = request.FILES.get(
                "file"
            )

            # =================================================
            # GET MESSAGE
            # =================================================

            if request.content_type.startswith(
                "multipart/form-data"
            ):

                message = request.POST.get(
                    "message",
                    ""
                ).strip()

            else:

                try:

                    data = json.loads(
                        request.body
                    )

                except json.JSONDecodeError:

                    return JsonResponse(
                        {
                            "error":
                            "Invalid message format."
                        },
                        status=400
                    )

                message = data.get(
                    "message",
                    ""
                ).strip()

            # =================================================
            # VALIDATE
            # =================================================

            if not message and not uploaded_file:

                return JsonResponse(
                    {
                        "error":
                        "Please enter a message or attach a file."
                    },
                    status=400
                )

            # =================================================
            # API KEY
            # =================================================

            api_key = os.getenv(
                "GEMINI_API_KEY"
            )

            if not api_key:

                return JsonResponse(
                    {
                        "error":
                        "Gemini API key not found. Check your .env file."
                    },
                    status=500
                )

            # =================================================
            # CLIENT
            # =================================================

            client = genai.Client(
                api_key=api_key
            )

            # =================================================
            # SYSTEM INSTRUCTION
            # =================================================

            system_instruction = (
                "You are IntelliGap AI, a student learning assistant. "

                "Answer clearly using simple language. "

                "Help students with Python, Java, AI, ML, "
                "DBMS, HTML, CSS, JavaScript, mathematics "
                "and academic subjects. "

                "When an image is provided, inspect it carefully. "

                "Read visible text, diagrams, tables, equations "
                "and other useful information from the image. "

                "When a document is provided, analyze its content "
                "and answer based on that content. "

                "If the student uploads a question, solve it. "

                "Give step-by-step explanations when appropriate. "

                "Do not pretend that you analyzed a file if "
                "the file was not actually received."
            )

            # =================================================
            # NORMAL TEXT CHAT
            # =================================================

            if not uploaded_file:

                ChatMessage.objects.create(
                    session_key=owner,
                    role="user",
                    message=message,
                )

                interaction = client.interactions.create(

                    model=GEMINI_MODEL,

                    system_instruction=
                    system_instruction,

                    generation_config={
                        "thinking_level": "low",
                        "max_output_tokens": 500,
                    },

                    input=message,
                )

            # =================================================
            # FILE CHAT
            # =================================================

            else:

                # ------------------------------------------------
                # FILE SIZE
                # ------------------------------------------------

                max_size = 50 * 1024 * 1024

                if uploaded_file.size > max_size:

                    return JsonResponse(
                        {
                            "error":
                            "File is too large. Maximum size is 50 MB."
                        },
                        status=400
                    )

                # ------------------------------------------------
                # CATEGORY
                # ------------------------------------------------

                category, mime_type = (
                    get_file_category(
                        uploaded_file
                    )
                )

                if category is None:

                    return JsonResponse(
                        {
                            "error":
                            (
                                "Unsupported file type. "
                                "Please use JPG, PNG, GIF, WEBP, "
                                "PDF, DOCX, PPTX, XLSX, TXT or CSV."
                            )
                        },
                        status=400
                    )

                # ------------------------------------------------
                # SAVE USER MESSAGE
                # ------------------------------------------------

                if message:

                    saved_message = (
                        f"{message}\n\n"
                        f"📎 {uploaded_file.name}"
                    )

                else:

                    saved_message = (
                        "Please analyze this file.\n\n"
                        f"📎 {uploaded_file.name}"
                    )

                ChatMessage.objects.create(
                    session_key=owner,
                    role="user",
                    message=saved_message,
                )

                # ------------------------------------------------
                # USER PROMPT
                # ------------------------------------------------

                if message:

                    user_prompt = message

                else:

                    user_prompt = (
                        "Analyze this uploaded file carefully. "
                        "Explain the important information clearly."
                    )

                # =================================================
                # IMAGE
                # =================================================

                if category == "image":

                    image_bytes = uploaded_file.read()

                    image_base64 = base64.b64encode(
                        image_bytes
                    ).decode(
                        "utf-8"
                    )

                    input_data = [

                        {
                            "type": "text",
                            "text": user_prompt
                        },

                        {
                            "type": "image",
                            "data": image_base64,
                            "mime_type": mime_type
                        }

                    ]

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 800,
                        },

                        input=input_data,
                    )

                # =================================================
                # PDF
                # =================================================

                elif category == "pdf":

                    gemini_file = upload_pdf_to_gemini(
                        client,
                        uploaded_file
                    )

                    input_data = [

                        {
                            "type": "text",
                            "text": user_prompt
                        },

                        {
                            "type": "document",
                            "uri": gemini_file.uri,
                            "mime_type": "application/pdf"
                        }

                    ]

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 1000,
                        },

                        input=input_data,
                    )

                # =================================================
                # TXT / CSV
                # =================================================

                elif category == "text":

                    file_text = extract_text_file(
                        uploaded_file
                    )

                    # Prevent unnecessarily huge prompts.
                    file_text = file_text[
                        :100000
                    ]

                    combined_prompt = (
                        f"{user_prompt}\n\n"
                        f"Here is the uploaded file content:\n\n"
                        f"{file_text}"
                    )

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 1000,
                        },

                        input=combined_prompt,
                    )

                # =================================================
                # DOCX
                # =================================================

                elif category == "docx":

                    document_text = extract_docx(
                        uploaded_file
                    )

                    document_text = document_text[
                        :100000
                    ]

                    combined_prompt = (
                        f"{user_prompt}\n\n"
                        f"Uploaded DOCX content:\n\n"
                        f"{document_text}"
                    )

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 1000,
                        },

                        input=combined_prompt,
                    )

                # =================================================
                # PPTX
                # =================================================

                elif category == "pptx":

                    presentation_text = extract_pptx(
                        uploaded_file
                    )

                    presentation_text = presentation_text[
                        :100000
                    ]

                    combined_prompt = (
                        f"{user_prompt}\n\n"
                        f"Uploaded PowerPoint content:\n\n"
                        f"{presentation_text}"
                    )

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 1000,
                        },

                        input=combined_prompt,
                    )

                # =================================================
                # XLSX
                # =================================================

                elif category == "xlsx":

                    spreadsheet_text = extract_xlsx(
                        uploaded_file
                    )

                    spreadsheet_text = spreadsheet_text[
                        :100000
                    ]

                    combined_prompt = (
                        f"{user_prompt}\n\n"
                        f"Uploaded Excel content:\n\n"
                        f"{spreadsheet_text}"
                    )

                    interaction = client.interactions.create(

                        model=GEMINI_MODEL,

                        system_instruction=
                        system_instruction,

                        generation_config={
                            "thinking_level": "low",
                            "max_output_tokens": 1000,
                        },

                        input=combined_prompt,
                    )

                else:

                    return JsonResponse(
                        {
                            "error":
                            "This file type is not currently supported."
                        },
                        status=400
                    )

            # =================================================
            # AI RESPONSE
            # =================================================

            reply = (
                interaction.output_text
                or
                "Sorry, I couldn't generate a response."
            )

            # =================================================
            # SAVE AI RESPONSE
            # =================================================

            ChatMessage.objects.create(
                session_key=owner,
                role="assistant",
                message=reply,
            )

            # =================================================
            # RETURN
            # =================================================

            return JsonResponse(
                {
                    "reply": reply
                }
            )

        # =====================================================
        # ERROR
        # =====================================================

        except Exception as e:

            print()
            print(
                "======================================"
            )

            print(
                "INTELLIGAP GEMINI ERROR"
            )

            print(
                type(e).__name__
            )

            print(
                str(e)
            )

            print(
                "======================================"
            )

            print()

            return JsonResponse(
                {
                    "error":
                    (
                        "File analysis failed. "
                        "Please check the Django terminal."
                    )
                },
                status=500
            )

    # ========================================================
    # INVALID METHOD
    # ========================================================

    return JsonResponse(
        {
            "error":
            "Invalid request method."
        },
        status=405
    )


# ============================================================
# KNOWLEDGE GAP
# ============================================================

def knowledge_gap(request):

    return render(
        request,
        "knowledge_gap.html"
    )